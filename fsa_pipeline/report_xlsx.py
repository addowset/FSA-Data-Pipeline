"""The customer-facing workbook.

This is the product. Everything else in the pipeline exists so that this file
can be attached to a reply, and it is the only evidence a supplier has that any
of the rest of it works.

Two tabs, deliberately. "New venues" is the list, laid out so it can be sorted
and filtered by someone who lives in a spreadsheet. "About this report" carries
the things that protect both sides: when the data was correct, what was taken
out before they saw it, and what the blank cells mean. A customer who finds a
discrepancy and has already been told how to read the file is a customer with a
question; one who has not been told is a customer who thinks they were sold
something careless.

No hygiene rating appears anywhere. The FSA's terms require a displayed rating
to be current or to carry the date it was accurate as of, and an out-of-date
one can be an offence. Nobody is buying the rating, so the obligation is not
worth taking on.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fsa_pipeline.csv_export import (
    companies_house_url,
    fhrs_url,
    google_maps_url,
    join_address,
)
from fsa_pipeline.report_filter import company_is_corroborated

FONT = "Calibri"
INK = "1F3330"
HEADER_FILL = PatternFill("solid", fgColor="1F3330")
TAKEOVER_FILL = PatternFill("solid", fgColor="F4F1E8")
RULE = Side(style="thin", color="D5DBD8")
BORDER = Border(left=RULE, right=RULE, top=RULE, bottom=RULE)
LINK_FONT = Font(name=FONT, size=10, color="1D5748", underline="single")

COLUMNS = [
    ("Venue", 30),
    ("Address", 44),
    ("Postcode", 11),
    ("Council", 22),
    ("Premises", 20),
    ("Status", 17),
    ("Previously", 26),
    ("First listed", 12),
    ("Company", 26),
    ("Company no.", 13),
    ("Incorporated", 13),
    ("Note", 30),
    ("Map", 7),
    ("FHRS record", 12),
]

PREMISES_LABELS = {
    "Restaurant/Cafe/Canteen": "Cafe / restaurant",
    "Pub/bar/nightclub": "Pub / bar",
    "Hotel/bed & breakfast/guest house": "Hotel / B&B",
    "Other catering premises": "Other catering",
    "Takeaway/sandwich shop": "Takeaway",
    "Retailers - other": "Food shop",
    "Retailers - supermarkets/hypermarkets": "Supermarket",
    "Mobile caterer": "Mobile caterer",
}


def premises_label(business_type: str | None) -> str:
    return PREMISES_LABELS.get(business_type or "", business_type or "")


SMALL_WORDS = {"and", "the", "of", "on", "in", "at", "to", "by", "for", "a", "an"}


def tidy_name(value: str | None) -> str:
    """Case-correct a name the council shouted.

    About one trading name in fifteen arrives entirely upper case -- "BLIND COW
    OXTON", "D'MUMBAI STATION KITCHEN AND BAR" -- because that is how the
    council's clerk typed it. Left alone it makes the whole file look scraped.
    Only fully upper-case names are touched: a name with any lower case in it is
    the owner's own styling and is left exactly as it is, so "munchbox" and
    "iTSU-style" spellings survive.
    """
    name = (value or "").strip()
    if not name or not name.isupper() or len(name) <= 3:
        return name
    words = name.title().split(" ")
    out = [words[0]] + [
        word.lower() if word.lower() in SMALL_WORDS else word for word in words[1:]
    ]
    return " ".join(out)


def build_report_row(record: dict) -> dict:
    """One record -> the cells a customer sees.

    Company columns are filled only where the match is corroborated by an exact
    registered-address hit. Of 102 company numbers checked by hand, four were
    the wrong company -- a dissolved Surrey entity attached to a Norfolk
    business, and similar. A customer cannot tell a wrong company name from a
    right one, so a blank is the honest cell.
    """
    address = join_address(
        record["address_line_1"], record["address_line_2"],
        record["address_line_3"], record["address_line_4"],
    )
    postcode = record["postcode_from_live_api"] or record["post_code"] or ""
    takeover = record["classification"] == "OPERATOR_CHANGE"
    corroborated = (
        bool(record["evidence_company_number"])
        and company_is_corroborated(record["reason"])
    )
    notes = []
    if record.get("batch_publication"):
        notes.append("Council published a batch this day, so it may have opened earlier")
    if takeover and not record["previous_business_name"]:
        notes.append("Previous occupier not recorded by the council")
    return {
        "Venue": tidy_name(record["business_name"]),
        "Address": address,
        "Postcode": postcode,
        "Council": record["authority_name"],
        "Premises": premises_label(record["business_type"]),
        "Status": "Change of hands" if takeover else "New opening",
        "Previously": tidy_name(record["previous_business_name"]) if takeover else "",
        "First listed": record["first_seen_date"],
        "Company": (record["company_name"] or "").strip() if corroborated else "",
        "Company no.": record["evidence_company_number"] if corroborated else "",
        "Incorporated": record["incorporation_date"] if corroborated else "",
        "Note": "; ".join(notes),
        "Map": google_maps_url(address, postcode),
        "FHRS record": fhrs_url(record["fhrsid"]),
        "_takeover": takeover,
        "_company_url": companies_house_url(record["evidence_company_number"]) if corroborated else "",
    }


def _header(ws, row=1):
    for index, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row, index, name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[row].height = 22


def write_report(path, rows, *, territory: str, authority_count: int,
                 since: str, until: str, audience: str, removed: dict,
                 created: dt.datetime | None = None, customer: str | None = None):
    """Write the workbook. Returns the stamp text, which the covering email
    should quote so the file and the message agree."""
    created = created or dt.datetime.now()
    stamp = "Correct at the time of report creation, %s" % created.strftime(
        "%d %B %Y at %H:%M"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "New venues"
    _header(ws)

    column_index = {name: i for i, (name, _w) in enumerate(COLUMNS, start=1)}
    company_column = column_index["Company"]

    for offset, row in enumerate(rows):
        r = offset + 2
        for index, (name, _width) in enumerate(COLUMNS, start=1):
            value = row[name]
            cell = ws.cell(r, index)
            if name == "Map" and value:
                cell.value = "Map"
                cell.hyperlink = value
                cell.font = LINK_FONT
            elif name == "FHRS record" and value:
                cell.value = "FHRS"
                cell.hyperlink = value
                cell.font = LINK_FONT
            else:
                cell.value = value
                cell.font = Font(name=FONT, size=10, color=INK)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top", wrap_text=name in ("Address", "Note")
            )
            if row["_takeover"]:
                cell.fill = TAKEOVER_FILL
        if row["_company_url"]:
            cell = ws.cell(r, company_column)
            cell.hyperlink = row["_company_url"]
            cell.font = LINK_FONT

    last = max(ws.max_row, 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), last)

    _write_notes(wb, stamp, territory, authority_count, since, until, audience,
                 removed, rows, customer)
    wb.save(path)
    return stamp


def _write_notes(wb, stamp, territory, authority_count, since, until, audience,
                 removed, rows, customer):
    ws = wb.create_sheet("About this report")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 86

    title = Font(name=FONT, bold=True, size=14, color=INK)
    label = Font(name=FONT, bold=True, size=10, color=INK)
    body = Font(name=FONT, size=10, color=INK)
    quiet = Font(name=FONT, size=10, italic=True, color="5B6A64")

    new = sum(1 for row in rows if not row["_takeover"])
    takeovers = len(rows) - new
    with_company = sum(1 for row in rows if row["Company"])

    ws["A1"] = "New food and drink venues"
    ws["A1"].font = title
    row = 3
    facts = [
        ("Prepared for", customer or ""),
        ("Area", "%s (%d local %s)" % (territory, authority_count,
                                       "authority" if authority_count == 1 else "authorities")),
        ("Period", "%s to %s" % (since, until)),
        ("Venues", "%d - %d new openings, %d changes of hands" % (len(rows), new, takeovers)),
        ("Accuracy", stamp),
    ]
    for name, value in facts:
        if not value:
            continue
        ws.cell(row, 1, name).font = label
        ws.cell(row, 2, value).font = body
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    ws.cell(row, 1, "What this is").font = label
    for line in [
        "Every business that serves food in the UK has to register with its local council "
        "before it opens. This is the week's new registrations for your area, taken from "
        "all 363 council registers, with the ones you cannot sell to removed.",
        "Just over half are changes of hands rather than new openings. Those are the rows "
        "shaded cream. The site was already trading, but whoever buys for it has just "
        "changed, which is usually the better moment to be in touch.",
    ]:
        ws.cell(row, 2, line).font = body
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 32
        row += 1

    row += 1
    ws.cell(row, 1, "What was taken out").font = label
    total = len(rows) + sum(
        count for reason, count in removed.items() if reason != "outside territory"
    )
    ws.cell(row, 2, "Your area produced %d register entries in this period. %d of them are in "
                    "your list. The rest were taken out before you saw it:"
            % (total, len(rows))).font = body
    ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 30
    row += 1
    removed_lines = {
        "unclassified": "register entries that were not a new opening or a change of hands",
        "venue type": "not the kind of premises this list covers",
        "incomplete address": "the council published no usable street address",
        "chain": "branches of national chains",
        "institution": "school, hospital, care home and prison kitchens",
        "re-issue": "the same business re-registered under a new reference, not a real change",
    }
    any_removed = False
    for reason, text in removed_lines.items():
        count = removed.get(reason, 0)
        if not count:
            continue
        ws.cell(row, 2, "%d %s" % (count, text)).font = body
        row += 1
        any_removed = True
    if not any_removed:
        ws.cell(row, 2, "Nothing in this area needed removing.").font = body
        row += 1

    row += 1
    ws.cell(row, 1, "The columns").font = label
    for line in [
        "Previously - the business that was at this address before, where the council "
        "recorded one. Blank on a new opening.",
        "Company, Company no., Incorporated - filled in only where the registered company "
        "address is an exact match for the venue address. Where the match was less certain "
        "these are left blank on purpose, because a plausible but wrong company is worse "
        "than none.",
        "First listed - the day the registration appeared on the council's register. Where "
        "a council publishes in batches the Note column says so, and the venue may have "
        "opened earlier.",
        "Map, FHRS record - links. The FHRS link opens the council's own entry for the venue.",
    ]:
        ws.cell(row, 2, line).font = body
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 32
        row += 1

    row += 1
    for line in [
        "Councils can amend or withdraw a registration after publishing it, so this file "
        "describes the register as it stood at the time above rather than making an "
        "open-ended claim. If something here does not match what you find, tell me and I "
        "will look at it.",
        "Contains public sector information licensed under the Open Government Licence v3.0. "
        "Opendo is not affiliated with or endorsed by the Food Standards Agency.",
        "Opendo - andrew@opendo.co.uk - opendo.co.uk",
    ]:
        ws.cell(row, 2, line).font = quiet
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1
