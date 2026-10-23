import datetime as dt

from openpyxl import load_workbook

from fsa_pipeline.report_xlsx import build_report_row, premises_label, tidy_name, write_report


def record(**overrides):
    base = {
        "fhrsid": 123456,
        "authority_name": "Bristol",
        "business_name": "The Cafe",
        "business_type": "Restaurant/Cafe/Canteen",
        "classification": "NEW_VENUE",
        "address_line_1": "1 High Street",
        "address_line_2": "Clifton",
        "address_line_3": None,
        "address_line_4": None,
        "post_code": "BS8 1AA",
        "postcode_from_live_api": None,
        "previous_business_name": None,
        "first_seen_date": "2026-10-20",
        "reason": "",
        "evidence_company_number": None,
        "company_name": None,
        "incorporation_date": None,
        "batch_publication": 0,
    }
    base.update(overrides)
    return base


# --- name tidying -----------------------------------------------------------

def test_shouted_names_are_case_corrected():
    assert tidy_name("BLIND COW OXTON") == "Blind Cow Oxton"


def test_small_words_stay_lower_case():
    assert tidy_name("FISH AND CHIPS OF LEEDS") == "Fish and Chips of Leeds"


def test_the_owners_own_styling_is_left_alone():
    # Any lower case at all means the owner chose it.
    assert tidy_name("munchbox") == "munchbox"
    assert tidy_name("Wok and fire") == "Wok and fire"


def test_short_names_are_left_alone():
    assert tidy_name("KFC") == "KFC"


# --- rows -------------------------------------------------------------------

def test_premises_labels_are_human():
    assert premises_label("Restaurant/Cafe/Canteen") == "Cafe / restaurant"
    assert premises_label("Pub/bar/nightclub") == "Pub / bar"


def test_new_opening_has_no_previous_occupier():
    row = build_report_row(record())
    assert row["Status"] == "New opening"
    assert row["Previously"] == ""


def test_takeover_names_the_previous_occupier():
    row = build_report_row(record(
        classification="OPERATOR_CHANGE", previous_business_name="Jack Dempsie's",
    ))
    assert row["Status"] == "Change of hands"
    assert row["Previously"] == "Jack Dempsie's"
    assert row["_takeover"] is True


def test_company_columns_need_address_corroboration():
    corroborated = build_report_row(record(
        evidence_company_number="12345678",
        company_name="THE CAFE LTD",
        incorporation_date="2026-08-01",
        reason="Registered address is an exact match for this establishment's address",
    ))
    assert corroborated["Company no."] == "12345678"
    assert corroborated["_company_url"].endswith("/12345678")


def test_a_weak_company_match_leaves_the_cells_blank():
    # A plausible but wrong company is worse than a blank, because the customer
    # cannot tell which it is.
    weak = build_report_row(record(
        evidence_company_number="12345678",
        company_name="SOME OTHER LTD",
        incorporation_date="2026-08-01",
        reason="Name similarity 0.91",
    ))
    assert weak["Company"] == ""
    assert weak["Company no."] == ""
    assert weak["Incorporated"] == ""
    assert weak["_company_url"] == ""


def test_batch_publication_is_explained_in_the_note():
    row = build_report_row(record(batch_publication=1))
    assert "earlier" in row["Note"]


# --- the workbook -----------------------------------------------------------

def test_workbook_has_both_tabs_and_the_accuracy_stamp(tmp_path):
    path = tmp_path / "report.xlsx"
    rows = [build_report_row(record())]
    stamp = write_report(
        path, rows, territory="Bristol", authority_count=1,
        since="2026-10-16", until="2026-10-22", audience="all",
        removed={"unclassified": 3}, created=dt.datetime(2026, 10, 23, 13, 32),
        customer="Test Roasters",
    )
    assert stamp == "Correct at the time of report creation, 23 October 2026 at 13:32"

    wb = load_workbook(path)
    assert wb.sheetnames == ["New venues", "About this report"]

    notes = "\n".join(
        str(cell.value) for row in wb["About this report"].iter_rows()
        for cell in row if cell.value
    )
    assert stamp in notes
    assert "Open Government Licence v3.0" in notes
    assert "Test Roasters" in notes
    # No hygiene rating anywhere: the FSA's terms make a stale one a liability
    # and nobody is buying it.
    assert "rating" not in notes.lower()


def test_an_empty_report_still_writes_a_readable_file(tmp_path):
    path = tmp_path / "empty.xlsx"
    write_report(
        path, [], territory="Cambridgeshire", authority_count=6,
        since="2026-10-16", until="2026-10-22", audience="roaster", removed={},
    )
    wb = load_workbook(path)
    assert wb["New venues"].max_row == 1          # header only
